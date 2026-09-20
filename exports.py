"""
Excel and Word report builders -- the same "research desk" palette as the app.

Both files are built from the scraped data only, and every piece of scraped
text is written as text: openpyxl stores any string starting with "=" as a
formula, which turned PDF text like "= 175 559 3,286" into formulas Excel
then "repaired" away (and would let a hostile page inject a working one).
"""

import io
import re
from datetime import datetime

import pandas as pd

from core import (DOCX_AVAILABLE, OPENPYXL_AVAILABLE, STATEMENT_TABLES, _count_downloaded,
                  _to_number, correlate_numbers)

if DOCX_AVAILABLE:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

if OPENPYXL_AVAILABLE:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink


# --------------------------------------------------------------------------
# Shared export helpers -- same "research desk" palette as the app theme
# --------------------------------------------------------------------------

INK, ACCENT, MUTED, LINE, PAPER, WHITE = "1C1F26", "1F4E79", "6B6F78", "E3E0D8", "F6F4EF", "FFFFFF"
GOOD, WARN, BAD, GOLD = "1F7A4D", "A8620A", "B4382F", "C08A3E"
TONE_COLOR = {"Positive": GOOD, "Neutral": WARN, "Mixed": WARN, "Cautious": BAD, "Negative": BAD}
TINT = {GOOD: "E5F2EA", WARN: "FBF1E1", BAD: "F7E6E4", MUTED: "EEEDEA"}
SERIF, SANS = "Georgia", "Calibri"  # both ship with Windows and macOS Office


def _parse_ratio(raw):
    """screener.in ratio text -> (number, Excel number format); text that isn't
    one number (e.g. High / Low "₹1,612/1,226") comes back as (text, None)."""
    s = str(raw or "").strip()
    num = None if "/" in s else _to_number(re.sub(r"₹|Cr\.?", "", s))
    if num is None:
        return s, None
    if s.endswith("%"):
        return num / 100, "0.00%"
    if "Cr" in s:
        return num, '"₹"#,##0" Cr"'
    if s.startswith("₹"):
        return num, '"₹"#,##0' if float(num).is_integer() else '"₹"#,##0.00'
    return num, "#,##0.0#"


def _quarterly_rows(qdf):
    """(quarter labels, [(metric, values, is_pct)]) with % rows as fractions."""
    if qdf is None or qdf.empty:
        return [], []
    quarters = [c for c in qdf.columns if c != "Metric"]
    rows = []
    for _, row in qdf.iterrows():
        metric = str(row["Metric"])
        if metric.lower().startswith("raw pdf"):
            continue
        is_pct = "%" in metric
        values = [_to_number(row[q]) for q in quarters]
        rows.append((metric, [v / 100 if is_pct and v is not None else v for v in values], is_pct))
    return quarters, rows


def _qoq(values, is_pct):
    """Latest-quarter change: percentage points for margin rows, % otherwise."""
    if len(values) < 2 or values[-1] is None or values[-2] is None:
        return None
    if is_pct:
        return (values[-1] - values[-2]) * 100
    return (values[-1] - values[-2]) / abs(values[-2]) if values[-2] else None


COST_ROWS = ("Expenses", "Interest", "Depreciation", "Tax")  # a rise here is bad news


def _qoq_color(metric, change):
    if not change:
        return MUTED
    return GOOD if (change > 0) != metric.startswith(COST_ROWS) else BAD


def _report_model(results):
    return next((r.get("ai_model") for r in results if r.get("ai_model")), None)


# --------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------

def _xl_put(ws, row, col, value=None, *, formula=False, size=10, bold=False, italic=False, color=INK,
            font=SANS, fill=None, h=None, v="center", wrap=False, fmt=None, border=None, link=None):
    """The ONLY way the exporter writes a cell. openpyxl stores any string that
    starts with "=" as a formula: scraped PDF text like "= 175 559 3,286" became
    invalid formulas that Excel "repaired" away (and a hostile page could inject
    working ones). Scraped text is therefore always stored as text; only the
    formulas this module builds itself pass formula=True."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("=") and not formula:
        cell.data_type = "s"
    cell.font = Font(name=font, size=size, bold=bold, italic=italic, color=color,
                     underline="single" if link else None)
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    if link and link.startswith("#"):  # jump within the workbook
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=link[1:])
    elif link:
        cell.hyperlink = link
    return cell


def _xl_merge(ws, row, col1, col2, value=None, **style):
    cell = _xl_put(ws, row, col1, value, **style)
    for c in range(col1 + 1, col2 + 1):  # keep fill/borders continuous across the merge
        _xl_put(ws, row, c, None, fill=style.get("fill"), border=style.get("border"))
    if col2 > col1:
        ws.merge_cells(start_row=row, start_column=col1, end_row=row, end_column=col2)
    return cell


def _xl_lines(text, chars_per_line):
    return max(1, sum(-(-max(len(part), 1) // chars_per_line) for part in str(text).split("\n")))


def _xl_section(ws, row, title, last_col, note=None):
    """Serif section title with an ink rule underneath, like the app's headings."""
    rule = Border(bottom=Side(style="medium", color=ACCENT))
    _xl_merge(ws, row, 2, last_col, title, font=SERIF, size=13, bold=True, border=rule, v="bottom")
    ws.row_dimensions[row].height = 24
    if note:
        _xl_put(ws, row + 1, 2, note, size=9, italic=True, color=MUTED)
        return row + 2
    return row + 1


def _xl_pill(ws, row, col1, col2, label):
    color = TONE_COLOR.get(label, MUTED)
    _xl_merge(ws, row, col1, col2, (label or "n/a").upper(), size=10, bold=True, color=color,
              fill=TINT[color], h="center")


def _xl_sheet_setup(ws, widths, tab_color):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _xl_ticker_sheet(wb, r):
    """One readable report per company. Returns cell refs the Overview links to."""
    from openpyxl.chart import BarChart, Reference, Series
    from openpyxl.formatting.rule import CellIsRule

    ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", r["ticker"])[:31] or "Ticker")
    quarters, qrows = _quarterly_rows(r.get("quarterly_df"))
    first_q = 4                                   # quarters start in column D
    last_col = max(17, first_q + len(quarters))   # grid runs B..Q at least
    qoq_col = first_q + len(quarters)
    _xl_sheet_setup(ws, {"A": 2, **{get_column_letter(c): 11.5 for c in range(2, last_col + 1)}},
                    TONE_COLOR.get(r.get("ai_verdict"), ACCENT if r["ok"] else BAD))
    chars_per_line = int((last_col - 1) * 11.5 * 1.15)
    thin = Border(bottom=Side(style="thin", color=LINE))
    refs = {"sheet": ws.title}

    _xl_merge(ws, 2, 2, last_col - 2, r["ticker"], font=SERIF, size=24, bold=True)
    ws.row_dimensions[2].height = 34
    _xl_merge(ws, 3, 2, last_col - 2, r.get("company_name") or "", size=11, color=MUTED)
    _xl_put(ws, 2, last_col, "← Overview", size=9, color=ACCENT, h="right", link="#'Overview'!A1")

    if not r["ok"]:
        msg = r.get("error_friendly") or r.get("error") or "Unknown error."
        _xl_merge(ws, 5, 2, last_col, f"Could not analyse this ticker: {msg}", color=BAD, fill=TINT[BAD], wrap=True)
        ws.row_dimensions[5].height = 15 * _xl_lines(msg, chars_per_line) + 10
        return refs

    row = 5
    if r.get("ai_verdict") or (r.get("news") or {}).get("sentiment"):
        _xl_put(ws, row, 2, "AI VERDICT", size=8, bold=True, color=MUTED)
        _xl_pill(ws, row, 3, 4, r.get("ai_verdict"))
        _xl_put(ws, row, 6, "NEWS SENTIMENT", size=8, bold=True, color=MUTED)
        _xl_pill(ws, row, 8, 9, (r.get("news") or {}).get("sentiment"))
        ws.row_dimensions[row].height = 20
        row += 2

    if r.get("ai_summary"):
        row = _xl_section(ws, row, "Analysis", last_col)
        _xl_merge(ws, row, 2, last_col, r["ai_summary"], size=11, fill=PAPER, wrap=True, v="top",
                  border=Border(left=Side(style="thick", color=ACCENT)))
        ws.row_dimensions[row].height = 16 * _xl_lines(r["ai_summary"], chars_per_line) + 12
        _xl_put(ws, row + 1, 2, f"Generated locally by {r.get('ai_model')} from screener.in data. Not investment advice.",
                size=8, italic=True, color=MUTED)
        row += 3

    if r.get("metrics"):
        row = _xl_section(ws, row, "Key ratios", last_col)
        items = list(r["metrics"].items())
        for i, (name, raw) in enumerate(items):
            tile_row, tile_col = row + (i // 5) * 3, 2 + (i % 5) * 3
            value, fmt = _parse_ratio(raw)
            _xl_merge(ws, tile_row, tile_col, tile_col + 2, name.upper(), size=8, bold=True, color=MUTED)
            _xl_merge(ws, tile_row + 1, tile_col, tile_col + 2, value, font=SERIF, size=15, h="left", fmt=fmt)
            ws.row_dimensions[tile_row + 1].height = 22
        row += -(-len(items) // 5) * 3 + 1

    pros, cons = (r.get("pros") or [])[:6], (r.get("cons") or [])[:6]
    if pros or cons:
        row = _xl_section(ws, row, "Strengths and risks", last_col)
        mid = 2 + (last_col - 1) // 2
        _xl_merge(ws, row, 2, mid - 1, "STRENGTHS", size=9, bold=True, color=GOOD, border=thin)
        _xl_merge(ws, row, mid + 1, last_col, "RISKS", size=9, bold=True, color=BAD, border=thin)
        row += 1
        width = int((mid - 2) * 11.5 * 1.1)
        for i in range(max(len(pros), len(cons))):
            p, c = (pros[i] if i < len(pros) else ""), (cons[i] if i < len(cons) else "")
            _xl_merge(ws, row, 2, mid - 1, f"•  {p}" if p else None, wrap=True, v="top")
            _xl_merge(ws, row, mid + 1, last_col, f"•  {c}" if c else None, wrap=True, v="top")
            ws.row_dimensions[row].height = 15 * max(_xl_lines(p, width), _xl_lines(c, width)) + 4
            row += 1
        row += 1

    news = r.get("news") or {}
    if news.get("headlines"):
        note = (f"Summarised locally by {news.get('model')} from the headlines below." if news.get("summary") else None)
        row = _xl_section(ws, row, "Recent news", last_col, note=note)
        if news.get("summary"):
            _xl_merge(ws, row, 2, last_col, news["summary"], size=10.5, fill=PAPER, wrap=True, v="top")
            ws.row_dimensions[row].height = 15 * _xl_lines(news["summary"], chars_per_line) + 10
            row += 1
        for h in news["headlines"]:
            _xl_put(ws, row, 2, h["published"].strftime("%d %b %Y") if h.get("published") else "",
                    size=9, color=MUTED, border=thin)
            _xl_merge(ws, row, 3, 4, h["source"], size=9, color=MUTED, border=thin)
            _xl_merge(ws, row, 5, last_col, h["title"], color=ACCENT, link=h.get("url") or None, border=thin)
            row += 1
        row += 1

    if qrows:
        row = _xl_section(ws, row, "Quarterly results", last_col,
                          note="₹ crore. Last column: change vs previous quarter (percentage points for margin rows).")
        _xl_merge(ws, row, 2, 3, "Metric", size=9, bold=True, color=WHITE, fill=ACCENT)
        for j, q in enumerate(quarters):
            _xl_put(ws, row, first_q + j, str(q), size=9, bold=True, color=WHITE, fill=ACCENT, h="right")
        _xl_put(ws, row, qoq_col, "QoQ", size=9, bold=True, color=WHITE, fill=ACCENT, h="right")
        header_row, data_start = row, row + 1
        prev_col, last_q_col = get_column_letter(qoq_col - 2), get_column_letter(qoq_col - 1)
        for i, (metric, values, is_pct) in enumerate(qrows):
            row += 1
            band = PAPER if i % 2 else None
            strong = metric.startswith(("Sales", "Net Profit"))
            _xl_merge(ws, row, 2, 3, metric, bold=strong, fill=band)
            fmt = "0%" if is_pct else ("#,##0.00" if metric.startswith("EPS") else "#,##0")
            for j, v in enumerate(values):
                _xl_put(ws, row, first_q + j, v, fmt=fmt, h="right", fill=band, bold=strong)
            if len(quarters) >= 2:
                f = (f"=IFERROR(({last_q_col}{row}-{prev_col}{row})*100,\"\")" if is_pct else
                     f"=IFERROR({last_q_col}{row}/{prev_col}{row}-1,\"\")")
                _xl_put(ws, row, qoq_col, f, formula=True, fill=band, h="right", bold=True,
                        fmt='+0.0" pp";-0.0" pp";0.0" pp"' if is_pct else "+0.0%;-0.0%;0.0%")
            if metric.startswith("Sales"):
                refs["sales_row"] = row
            elif metric.startswith("Net Profit"):
                refs["profit_row"] = row
        qoq_letter = get_column_letter(qoq_col)
        for is_cost in (False, True):  # costs rising is red, everything else rising is green
            cells = " ".join(f"{qoq_letter}{data_start + i}" for i, (m, _, _) in enumerate(qrows)
                             if m.startswith(COST_ROWS) == is_cost)
            if cells:
                up, down = (BAD, GOOD) if is_cost else (GOOD, BAD)
                ws.conditional_formatting.add(cells, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=up)))
                ws.conditional_formatting.add(cells, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=down)))
        refs.update(qoq_col=get_column_letter(qoq_col), last_q_col=last_q_col)
        row += 2

        if "sales_row" in refs and len(quarters) >= 2:
            chart = BarChart()
            chart.type, chart.grouping, chart.overlap, chart.gapWidth = "col", "clustered", -10, 60
            chart.title = "Sales vs net profit (₹ Cr)"
            for src_row, title, color in [(refs["sales_row"], "Sales", ACCENT), (refs.get("profit_row"), "Net profit", GOLD)]:
                if src_row:
                    series = Series(Reference(ws, min_col=first_q, max_col=qoq_col - 1, min_row=src_row), title=title)
                    series.graphicalProperties.solidFill = color
                    series.graphicalProperties.line.solidFill = color
                    chart.series.append(series)
            chart.set_categories(Reference(ws, min_col=first_q, max_col=qoq_col - 1, min_row=header_row))
            chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl 3.1 hides axes by default
            chart.y_axis.numFmt, chart.y_axis.majorGridlines = "#,##0", None
            chart.legend.position = "b"
            chart.height, chart.width = 7.5, 2.3 * (last_col - 1)
            ws.add_chart(chart, f"B{row}")
            row += 17

    downloaded = r.get("downloaded") or {}
    if downloaded:
        row = _xl_section(ws, row, "Filings", last_col)
        for cat, infos in downloaded.items():
            for info in infos:
                ok = bool(info.get("path"))
                _xl_merge(ws, row, 2, 4, cat, border=thin)
                _xl_merge(ws, row, 5, 8, "Downloaded" if ok else str(info.get("error") or "Skipped"),
                          color=GOOD if ok else MUTED, border=thin, size=9)
                _xl_merge(ws, row, 9, last_col, info.get("url") or "", size=9, color=ACCENT,
                          link=info.get("url") or None, border=thin)
                row += 1
        row += 1

    figs = r.get("figures") or []
    if figs:
        _xl_put(ws, row, 2, f"{len(figs):,} figures were extracted from this company's filings. "
                            "Filter them on the Figures sheet.", size=9, italic=True, color=MUTED)
        _xl_put(ws, row + 1, 2, "→ Open Figures", size=9, color=ACCENT, link="#'Figures'!A1")
    return refs


def _xl_correlated_sheet(wb, results):
    """Every metric the app holds, with each source of it side by side --
    the detail behind the one-line summaries on the company sheets."""
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows = []
    for r in results:
        if not r.get("ok"):
            continue
        df = r.get("correlated")
        if df is None:
            df = correlate_numbers(r)
        for _, row in df.iterrows():
            rows.append((r["ticker"], row))
    if not rows:
        return
    ws = wb.create_sheet("All numbers")
    cols = [("Ticker", 12), ("Metric", 30), ("Group", 18), ("Screener", 15), ("Latest quarter", 15),
            ("Previous quarter", 16), ("QoQ", 10), ("In filings", 10), ("Filing latest", 14),
            ("Filing low", 13), ("Filing median", 14), ("Filing high", 13), ("Unit", 8), ("Sources", 34), ("Pages", 18)]
    _xl_sheet_setup(ws, {get_column_letter(i + 1): w for i, (_, w) in enumerate(cols)}, ACCENT)
    _xl_put(ws, 1, 1, "All numbers, correlated", font=SERIF, size=16, bold=True)
    _xl_put(ws, 2, 1, "One row per metric. 'Screener' is the published ratio, the quarter columns come from the "
                      "results table, and the filing columns summarise every mention of that metric in the PDFs -- "
                      "a wide low-to-high spread means the label caught more than one kind of number.",
            size=9, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 26
    for j, (name, _) in enumerate(cols, start=1):
        _xl_put(ws, 4, j, name, bold=True, color=WHITE, fill=ACCENT, wrap=True, h="center")
    for i, (ticker, row) in enumerate(rows, start=5):
        band = PAPER if i % 2 else None
        values = [ticker, row["Metric"], row["Group"], row["Screener"], row["Latest quarter"],
                  row["Previous quarter"], row["QoQ"], row["In filings"] or None, row["Filing latest"],
                  row["Filing low"], row["Filing median"], row["Filing high"], row["Unit"],
                  row["Sources"], row.get("Pages", "")]
        for j, v in enumerate(values, start=1):
            fmt = "0.0%" if j == 7 else ("#,##0.##" if j in (5, 6, 9, 10, 11, 12) else None)
            color = INK
            if j == 7 and isinstance(v, (int, float)):
                color = GOOD if v >= 0 else BAD
            _xl_put(ws, i, j, None if pd.isna(v) else v, size=9.5, fmt=fmt, fill=band, color=color,
                    bold=(j == 7 and isinstance(v, (int, float))))
    table = Table(displayName="CorrelatedNumbers", ref=f"A4:{get_column_letter(len(cols))}{4 + len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = table.ref


def _xl_figures_sheet(wb, results):
    """Every extracted number, all tickers, as one filterable Excel Table."""
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows = [(r["ticker"], f) for r in results for f in (r.get("figures") or [])]
    if not rows:
        return
    ws = wb.create_sheet("Figures")
    cols = [("Ticker", 12), ("Source", 20), ("Category", 18), ("Label", 32), ("Value", 14),
            ("Unit", 8), ("Currency", 9), ("Page", 7), ("Context", 90)]
    _xl_sheet_setup(ws, {get_column_letter(i + 1): w for i, (_, w) in enumerate(cols)}, MUTED)
    _xl_put(ws, 1, 1, "Figures extracted from filings", font=SERIF, size=16, bold=True)
    _xl_put(ws, 2, 1, "Found by pattern-matching the downloaded PDFs, not by AI. Labels are best-effort: "
                      "check the Context column before relying on a number.", size=9, italic=True, color=MUTED)
    ws.row_dimensions[1].height = 26
    for j, (name, _) in enumerate(cols, start=1):
        _xl_put(ws, 4, j, name, bold=True, color=WHITE, fill=ACCENT)
    for i, (ticker, f) in enumerate(rows, start=5):
        values = [ticker, f.get("Source"), f.get("Category"), f.get("Label"), f.get("Value"),
                  f.get("Unit"), f.get("Currency"), f.get("Page"), f.get("Context")]
        for j, v in enumerate(values, start=1):
            _xl_put(ws, i, j, v, fmt="#,##0.##" if j == 5 else None, size=9.5,
                    color=MUTED if j == 9 else INK)
    table = Table(displayName="ExtractedFigures", ref=f"A4:{get_column_letter(len(cols))}{4 + len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B5"


def build_excel_workbook(results: list[dict]) -> bytes:
    if not OPENPYXL_AVAILABLE:
        buf = io.BytesIO()
        pd.DataFrame([{"note": "Install openpyxl for a formatted workbook."}]).to_excel(buf, index=False)
        return buf.getvalue()
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule

    wb = openpyxl.Workbook()
    ov = wb.active
    ov.title = "Overview"
    refs = {r["ticker"]: _xl_ticker_sheet(wb, r) for r in results}
    _xl_correlated_sheet(wb, results)
    _xl_figures_sheet(wb, results)

    columns = [("Ticker", 11), ("Company", 30), ("AI verdict", 12), ("News", 11), ("Price", 11),
               ("Market cap", 16), ("P/E", 8), ("ROE", 9), ("ROCE", 9), ("Div. yield", 10),
               ("Sales, latest qtr", 16), ("Sales QoQ", 11), ("Profit QoQ", 11), ("Filings", 9), ("Figures", 10)]
    _xl_sheet_setup(ov, {"A": 2, **{get_column_letter(i + 2): w for i, (_, w) in enumerate(columns)}}, ACCENT)
    last_col = len(columns) + 1

    model = _report_model(results)
    _xl_put(ov, 2, 2, "Indian Stock Data Hub", font=SERIF, size=24, bold=True)
    ov.row_dimensions[2].height = 36
    _xl_put(ov, 3, 2, f"Equity research report · {datetime.now():%d %b %Y, %H:%M}"
                      + (f" · AI analysis by {model}, running locally" if model else ""), size=10, color=MUTED)

    head, first, last = 9, 10, 9 + len(results)
    col = {name: i + 2 for i, (name, _) in enumerate(columns)}
    L = {name: get_column_letter(c) for name, c in col.items()}
    rng = lambda name: f"{L[name]}{first}:{L[name]}{last}"  # noqa: E731
    tiles = [("Companies", f"=COUNTA({rng('Ticker')})"),
             ("Positive verdicts", f'=COUNTIF({rng("AI verdict")},"Positive")'),
             ("Cautious verdicts", f'=COUNTIF({rng("AI verdict")},"Cautious")'),
             ("Filings downloaded", f"=SUM({rng('Filings')})"),
             ("Figures extracted", f"=SUM({rng('Figures')})")]
    tile_spans = [(2, 3), (4, 6), (7, 9), (10, 12), (13, 16)]  # ~equal widths
    for (label, formula), (c1, c2) in zip(tiles, tile_spans):
        edge = Border(left=Side(style="thick", color=ACCENT))
        _xl_merge(ov, 5, c1, c2, label.upper(), size=8, bold=True, color=MUTED, fill=PAPER, border=edge)
        _xl_merge(ov, 6, c1, c2, formula, formula=True, font=SERIF, size=20, fill=PAPER, border=edge,
                  h="left", fmt="#,##0")
    ov.row_dimensions[6].height = 32

    for name, c in col.items():
        _xl_put(ov, head, c, name, size=9, bold=True, color=WHITE, fill=ACCENT, wrap=True,
                h="left" if name in ("Ticker", "Company") else "center")
    ov.row_dimensions[head].height = 30
    thin = Border(bottom=Side(style="thin", color=LINE))
    for i, r in enumerate(results):
        row, ref = first + i, refs[r["ticker"]]
        put = lambda name, value, **kw: _xl_put(ov, row, col[name], value, border=thin, **kw)  # noqa: E731
        put("Ticker", r["ticker"], bold=True, color=ACCENT, link=f"#'{ref['sheet']}'!A1")
        put("Company", r.get("company_name") or "", wrap=True)
        if not r["ok"]:
            _xl_merge(ov, row, col["AI verdict"], last_col, "Failed: " + (r.get("error_friendly") or r.get("error") or ""),
                      color=BAD, border=thin, size=9)
            continue
        verdict = r.get("ai_verdict")
        tone = TONE_COLOR.get(verdict, MUTED)
        put("AI verdict", verdict or "–", bold=True, color=tone, fill=TINT[tone] if verdict else None, h="center")
        sentiment = (r.get("news") or {}).get("sentiment")
        put("News", sentiment or "–", color=TONE_COLOR.get(sentiment, MUTED), h="center")
        m = r.get("metrics") or {}
        for name, key in [("Price", "Current Price"), ("Market cap", "Market Cap"), ("P/E", "Stock P/E"),
                          ("ROE", "ROE"), ("ROCE", "ROCE"), ("Div. yield", "Dividend Yield")]:
            value, fmt = _parse_ratio(m.get(key, ""))
            put(name, value if value != "" else "–", fmt=fmt, h="right")
        # Live links into the company sheet, so edits there flow through here.
        sheet = f"'{ref['sheet']}'"
        if ref.get("sales_row"):
            put("Sales, latest qtr", f"={sheet}!{ref['last_q_col']}{ref['sales_row']}", formula=True,
                fmt='"₹"#,##0" Cr"', h="right")
            put("Sales QoQ", f"={sheet}!{ref['qoq_col']}{ref['sales_row']}", formula=True,
                fmt="+0.0%;-0.0%;0.0%", h="right", bold=True)
        if ref.get("profit_row"):
            put("Profit QoQ", f"={sheet}!{ref['qoq_col']}{ref['profit_row']}", formula=True,
                fmt="+0.0%;-0.0%;0.0%", h="right", bold=True)
        put("Filings", _count_downloaded(r), h="right")
        put("Figures", len(r.get("figures") or []), fmt="#,##0", h="right")
        ov.row_dimensions[row].height = 20
    for name in ("Sales QoQ", "Profit QoQ"):
        ov.conditional_formatting.add(rng(name), CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=GOOD)))
        ov.conditional_formatting.add(rng(name), CellIsRule(operator="lessThan", formula=["0"], font=Font(color=BAD)))
    ov.freeze_panes = ov.cell(row=first, column=col["AI verdict"])
    _xl_put(ov, last + 2, 2, "Click a ticker to open its report. Verdicts are written by a local AI model from "
                             "scraped data and can be wrong; this is not investment advice.",
            size=8.5, italic=True, color=MUTED)

    wb.calculation.fullCalcOnLoad = True  # openpyxl can't store results; make Excel compute on open
    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    # Round-trip verification: never hand the user a file we can't re-open
    # ourselves. If openpyxl can't parse its own output, something upstream
    # (bad character, bad range, etc.) is wrong -- fail loudly here instead.
    check_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    for name in check_wb.sheetnames:
        for _ in check_wb[name].iter_rows():
            pass
    check_wb.close()
    return data


# --------------------------------------------------------------------------
# Word export
# --------------------------------------------------------------------------

def _w_xml(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el


# Child order of the property elements we touch, per the WordprocessingML
# schema. Word reports out-of-order properties as "unreadable content", so
# every property this module adds goes through _w_set rather than .append().
_W_ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd "
           "fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
           "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
             "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcBorders": "top start left bottom end right insideH insideV tl2br tr2bl",
    "tcMar": "top start left bottom end right",
    "pBdr": "top left bottom right between bar",
}
_W_ORDER = {k: [qn(f"w:{n}") for n in v.split()] for k, v in _W_ORDER.items()}


def _w_set(parent, child):
    """Put child into a w:*Pr-style element at its schema position, replacing
    any existing element of the same kind."""
    order = _W_ORDER[parent.tag.rsplit("}", 1)[-1]]
    for old in parent.findall(child.tag):
        parent.remove(old)
    rank = order.index(child.tag)
    for sibling in parent:
        if sibling.tag in order and order.index(sibling.tag) > rank:
            sibling.addprevious(child)
            return child
    parent.append(child)
    return child


def _w_border(parent_tag, edges):
    """edges: {edge: (size_eighths_pt, color, space_pt)} -> ordered w:tcBorders / w:pBdr."""
    box = _w_xml(parent_tag)
    for edge, (size, color, space) in edges.items():
        _w_set(box, _w_xml(f"w:{edge}", **{"w:val": "single", "w:sz": size, "w:space": space, "w:color": color}))
    return box


def _w_font(target, name):
    """Set a font on a style or run, dropping theme-font attributes that
    otherwise override the name (python-docx's template uses theme fonts)."""
    target.font.name = name
    rfonts = target.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:eastAsia"), name)


def _w_run(par, text, size=None, bold=False, italic=False, color=INK, font=None, caps=False, spacing=None):
    run = par.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    if font:
        _w_font(run, font)
    if caps:
        run.font.all_caps = True
    if spacing:
        _w_set(run.element.get_or_add_rPr(), _w_xml("w:spacing", **{"w:val": spacing}))
    return run


def _w_par(container, text="", size=None, color=INK, bold=False, italic=False, align=None,
           before=0, after=4, font=None, caps=False, spacing=None, keep_next=False):
    par = container.add_paragraph()
    fmt = par.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.keep_with_next = keep_next
    if align:
        par.alignment = align
    if text:
        _w_run(par, text, size, bold, italic, color, font, caps, spacing)
    return par


def _w_pill(par, label):
    color = TONE_COLOR.get(label, MUTED)
    run = _w_run(par, f"  {(label or 'n/a').upper()}  ", size=8.5, bold=True, color=color, spacing=10)
    _w_set(run.element.get_or_add_rPr(), _w_xml("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": TINT[color]}))


def _w_cell(cell, fill=None, borders=None, width=None, valign="center", margins=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        _w_set(tc_pr, _w_xml("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
    if borders:
        _w_set(tc_pr, _w_border("w:tcBorders", {e: (sz, c, 0) for e, (sz, c) in borders.items()}))
    if margins:
        mar = _w_xml("w:tcMar")
        for edge, twips in margins.items():
            _w_set(mar, _w_xml(f"w:{edge}", **{"w:w": twips, "w:type": "dxa"}))
        _w_set(tc_pr, mar)
    if width:
        cell.width = width
    cell.vertical_alignment = {"center": WD_ALIGN_VERTICAL.CENTER, "top": WD_ALIGN_VERTICAL.TOP}[valign]
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return cell


def _w_table(doc, rows, widths):
    """Fixed-layout table. python-docx only sets per-cell widths; Word, Pages and
    Quick Look size columns from the table grid, so set that and the layout too."""
    table = doc.add_table(rows=rows, cols=len(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    _w_set(tbl_pr, _w_xml("w:tblW", **{"w:w": int(sum(w.twips for w in widths)), "w:type": "dxa"}))
    _w_set(tbl_pr, _w_xml("w:tblLayout", **{"w:type": "fixed"}))
    for grid_col, w in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.w = w
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    return table


def _w_cell_text(cell, text, size=9.5, bold=False, color=INK, align=None, italic=False):
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    if align:
        par.alignment = align
    _w_run(par, str(text), size=size, bold=bold, italic=italic, color=color)
    return par


def _w_hyperlink(par, url, text, size=9.5, color=ACCENT):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = _w_xml("w:hyperlink", **{"r:id": r_id})
    run = _w_xml("w:r")
    rpr = _w_xml("w:rPr")
    rpr.append(_w_xml("w:color", **{"w:val": color}))
    rpr.append(_w_xml("w:sz", **{"w:val": int(size * 2)}))
    run.append(rpr)
    t = _w_xml("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def _w_heading(doc, text, level=2, before=14):
    par = doc.add_heading(text, level=level)
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(before), Pt(6)
    if level == 2:  # thin ink rule under section headings, echoing the app
        _w_set(par._p.get_or_add_pPr(), _w_border("w:pBdr", {"bottom": (6, ACCENT, 3)}))
    return par


def _w_setup(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    normal = doc.styles["Normal"]
    _w_font(normal, SANS)
    normal.font.size, normal.font.color.rgb = Pt(10.5), RGBColor.from_string(INK)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(4), 1.15
    for level, size in [(1, 24), (2, 13)]:
        style = doc.styles[f"Heading {level}"]
        _w_font(style, SERIF)
        style.font.size, style.font.bold = Pt(size), True
        style.font.color.rgb = RGBColor.from_string(INK)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
    _w_run(footer, f"Indian Stock Data Hub · {datetime.now():%d %b %Y}\tPage ", size=8, color=MUTED)
    run = _w_run(footer, "", size=8, color=MUTED)
    for kind, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if kind:
            run._r.append(_w_xml("w:fldChar", **{"w:fldCharType": kind}))
        else:
            instr = _w_xml("w:instrText")
            instr.text = text
            instr.set(qn("xml:space"), "preserve")
            run._r.append(instr)


def _w_company(doc, r):
    eyebrow = _w_par(doc, "Company report", size=8, color=MUTED, caps=True, spacing=30, after=0)
    eyebrow.paragraph_format.page_break_before = True  # no empty break paragraph between companies
    doc.add_heading(r["ticker"], level=1).paragraph_format.space_after = Pt(0)
    _w_par(doc, r.get("company_name") or "", size=12, color=MUTED, after=8)
    if not r["ok"]:
        _w_par(doc, f"Could not analyse this ticker: {r.get('error_friendly') or r.get('error')}", color=BAD)
        return

    news = r.get("news") or {}
    if r.get("ai_verdict") or news.get("sentiment"):
        par = _w_par(doc, after=10)
        _w_run(par, "AI VERDICT  ", size=8, bold=True, color=MUTED)
        _w_pill(par, r.get("ai_verdict"))
        _w_run(par, "      NEWS SENTIMENT  ", size=8, bold=True, color=MUTED)
        _w_pill(par, news.get("sentiment"))

    if r.get("ai_summary"):
        box = _w_table(doc, 1, [Cm(17)])
        cell = _w_cell(box.cell(0, 0), fill=PAPER, borders={"left": (24, ACCENT)},
                       margins={"top": 140, "bottom": 140, "left": 220, "right": 220})
        _w_cell_text(cell, r["ai_summary"], size=10.5)
        _w_par(doc, f"Generated locally by {r.get('ai_model')} from screener.in data.", size=8, italic=True,
               color=MUTED, before=3, after=6)

    metrics = list((r.get("metrics") or {}).items())
    if metrics:
        _w_heading(doc, "Key ratios")
        per_row = 5
        grid = _w_table(doc, -(-len(metrics) // per_row) * 2, [Cm(17 / per_row)] * per_row)
        for i, (name, value) in enumerate(metrics):
            label_cell = _w_cell(grid.cell((i // per_row) * 2, i % per_row))
            value_cell = _w_cell(grid.cell((i // per_row) * 2 + 1, i % per_row),
                                 borders={"bottom": (4, LINE)} if i // per_row == 0 else None)
            _w_cell_text(label_cell, name.upper(), size=7.5, bold=True, color=MUTED)
            par = value_cell.paragraphs[0]
            _w_run(par, value, size=13, font=SERIF)
            par.paragraph_format.space_after = Pt(8)

    pros, cons = (r.get("pros") or [])[:6], (r.get("cons") or [])[:6]
    if pros or cons:
        _w_heading(doc, "Strengths and risks")
        pc = _w_table(doc, 1, [Cm(8.3), Cm(0.4), Cm(8.3)])
        for idx, (title, items, color) in [(0, ("Strengths", pros, GOOD)), (2, ("Risks", cons, BAD))]:
            cell = _w_cell(pc.cell(0, idx), fill=TINT[color], valign="top",
                           margins={"top": 120, "bottom": 120, "left": 160, "right": 160})
            _w_cell_text(cell, title.upper(), size=8, bold=True, color=color)
            for item in items or ["None listed"]:
                par = cell.add_paragraph()
                par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(4), Pt(0)
                par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.35), Cm(-0.35)
                _w_run(par, f"•  {item}", size=9.5)

    quarters, qrows = _quarterly_rows(r.get("quarterly_df"))
    if qrows:
        shown = quarters[-5:]  # 13 quarters are unreadable in portrait; Excel has all of them
        _w_heading(doc, "Quarterly results")
        _w_par(doc, f"₹ crore, last {len(shown)} quarters. Full history in the Excel workbook.", size=8,
               italic=True, color=MUTED, after=4)
        widths = [Cm(4.2)] + [Cm(2.2)] * len(shown) + [Cm(1.8)]
        qt = _w_table(doc, len(qrows) + 1, widths)
        for j, text in enumerate(["Metric", *map(str, shown), "QoQ"]):
            cell = _w_cell(qt.cell(0, j), fill=ACCENT)
            _w_cell_text(cell, text, size=8.5, bold=True, color=WHITE,
                         align=None if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
        for i, (metric, values, is_pct) in enumerate(qrows, start=1):
            band = PAPER if i % 2 == 0 else None
            strong = metric.startswith(("Sales", "Net Profit"))
            _w_cell_text(_w_cell(qt.cell(i, 0), fill=band), metric, size=9, bold=strong)
            for j, v in enumerate(values[-len(shown):], start=1):
                text = "–" if v is None else (f"{v:.0%}" if is_pct else f"{v:,.2f}" if metric.startswith("EPS") else f"{v:,.0f}")
                _w_cell_text(_w_cell(qt.cell(i, j), fill=band), text, size=9, bold=strong, align=WD_ALIGN_PARAGRAPH.RIGHT)
            change = _qoq(values, is_pct)
            text = "–" if change is None else (f"{change:+.1f} pp" if is_pct else f"{change:+.1%}")
            color = _qoq_color(metric, change)
            _w_cell_text(_w_cell(qt.cell(i, len(shown) + 1), fill=band), text, size=9, bold=True, color=color,
                         align=WD_ALIGN_PARAGRAPH.RIGHT)

    _w_correlated(doc, r)

    if news.get("headlines"):
        _w_heading(doc, "Recent news")
        if news.get("summary"):
            _w_par(doc, news["summary"], after=2)
            _w_par(doc, f"Summarised locally by {news.get('model')} from the headlines below.", size=8,
                   italic=True, color=MUTED, after=6)
        for h in news["headlines"]:
            par = _w_par(doc, after=3)
            par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.35), Cm(-0.35)
            _w_run(par, "›  ", size=9.5, color=ACCENT, bold=True)
            if h.get("url"):
                _w_hyperlink(par, h["url"], h["title"])
            else:
                _w_run(par, h["title"], size=9.5)
            date = f" · {h['published']:%d %b}" if h.get("published") else ""
            _w_run(par, f"   {h['source']}{date}", size=8.5, color=MUTED)

    downloaded = r.get("downloaded") or {}
    figs = r.get("figures") or []
    if downloaded or figs:
        _w_heading(doc, "Filings")
        for cat, infos in downloaded.items():
            for info in infos:
                par = _w_par(doc, after=2)
                _w_run(par, f"{cat}   ", size=9.5, bold=True)
                if info.get("path"):
                    _w_run(par, "Downloaded   ", size=9, color=GOOD)
                    if info.get("url"):
                        _w_hyperlink(par, info["url"], "Open source ↗", size=9)
                else:
                    _w_run(par, str(info.get("error") or "Skipped"), size=9, color=MUTED)
        if figs:
            by_cat = pd.Series([f.get("Category") or "Other" for f in figs]).value_counts()
            _w_par(doc, f"{len(figs):,} figures extracted from these filings: "
                        + ", ".join(f"{cat} {n:,}" for cat, n in by_cat.items())
                        + ". The full, filterable list is on the Figures sheet of the Excel workbook.",
                   size=8.5, italic=True, color=MUTED, before=6)


def _w_correlated(doc, r, limit=14):
    """The best-corroborated metrics, with each source beside the others."""
    df = r.get("correlated")
    if df is None:
        df = correlate_numbers(r)
    if df is None or df.empty:
        return
    rows = df.head(limit)
    _w_heading(doc, "Key numbers, cross-checked")
    _w_par(doc, "Each metric as screener.in publishes it, as the last two quarters reported it, and as the "
                "downloaded filings state it. Where they disagree, trust the filing and check its page.",
           size=8.5, italic=True, color=MUTED, after=6)
    cols = [("Metric", 5.0), ("Screener", 2.6), ("Latest qtr", 2.4), ("Prev qtr", 2.4),
            ("QoQ", 1.9), ("In filings", 1.8), ("Filing range", 3.0)]
    table = _w_table(doc, len(rows) + 1, [Cm(w) for _, w in cols])
    for j, (name, _) in enumerate(cols):
        _w_cell_text(_w_cell(table.cell(0, j), fill=ACCENT), name, size=8.5, bold=True, color=WHITE,
                     align=WD_ALIGN_PARAGRAPH.CENTER if j else None)

    def num(value, pct=False):
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return "–"
        if pct:
            return f"{value:+.1%}"
        return f"{value:,.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else f"{value:,}"

    for i, (_, row) in enumerate(rows.iterrows(), start=1):
        band = PAPER if i % 2 else None
        spread = "–"
        if row["In filings"]:
            spread = f"{num(row['Filing low'])} – {num(row['Filing high'])}"
        cells = [str(row["Metric"])[:42], str(row["Screener"] or "–"), num(row["Latest quarter"]),
                 num(row["Previous quarter"]), num(row["QoQ"], pct=True),
                 str(row["In filings"] or "–"), spread]
        for j, text in enumerate(cells):
            color = INK
            if j == 4 and row["QoQ"] is not None and not pd.isna(row["QoQ"]):
                color = _qoq_color(str(row["Metric"]), row["QoQ"] * 100)
            _w_cell_text(_w_cell(table.cell(i, j), fill=band), text, size=8.5, color=color,
                         bold=(j == 4 and color != INK),
                         align=WD_ALIGN_PARAGRAPH.RIGHT if j else None)


def build_word_document(results: list[dict]) -> bytes:
    doc = Document()
    _w_setup(doc)
    model = _report_model(results)

    _w_par(doc, "Equity research report", size=9, color=MUTED, caps=True, spacing=40, before=60, after=4)
    _w_par(doc, "Indian Stock Data Hub", size=30, font=SERIF, bold=True, after=2)
    tickers = ", ".join(r["ticker"] for r in results)
    rule = _w_par(doc, f"{tickers} · {datetime.now():%d %B %Y}"
                       + (f" · analysis by {model}, run locally" if model else ""), size=10.5, color=MUTED, after=18)
    _w_set(rule._p.get_or_add_pPr(), _w_border("w:pBdr", {"bottom": (18, ACCENT, 10)}))

    _w_heading(doc, "At a glance", before=6)
    cols = [("Ticker", 2.0), ("Company", 5.2), ("AI verdict", 2.3), ("News", 2.0), ("Price", 1.9), ("P/E", 1.5), ("ROE", 2.1)]
    glance = _w_table(doc, len(results) + 1, [Cm(w) for _, w in cols])
    for j, (name, _) in enumerate(cols):
        _w_cell_text(_w_cell(glance.cell(0, j), fill=ACCENT), name, size=8.5, bold=True, color=WHITE,
                     align=WD_ALIGN_PARAGRAPH.RIGHT if j >= 4 else None)
    for i, r in enumerate(results, start=1):
        line = {"bottom": (4, LINE)}
        cells = [_w_cell(glance.cell(i, j), borders=line) for j in range(len(cols))]
        for c in cells:
            c.paragraphs[0].paragraph_format.space_before = Pt(3)
            c.paragraphs[0].paragraph_format.space_after = Pt(3)
        _w_cell_text(cells[0], r["ticker"], bold=True)
        _w_cell_text(cells[1], r.get("company_name") or "", size=9)
        if not r["ok"]:
            _w_cell_text(cells[2], "Failed", bold=True, color=BAD)
            continue
        _w_pill(cells[2].paragraphs[0], r.get("ai_verdict"))
        sentiment = (r.get("news") or {}).get("sentiment")
        _w_cell_text(cells[3], sentiment or "–", size=9, color=TONE_COLOR.get(sentiment, MUTED))
        m = r.get("metrics") or {}
        for c, key in zip(cells[4:], ["Current Price", "Stock P/E", "ROE"]):
            _w_cell_text(c, m.get(key, "–"), size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)

    _w_par(doc, "How this report was made", size=8, bold=True, color=MUTED, caps=True, before=18, after=2)
    _w_par(doc, "Ratios, quarterly results, pros and cons come from screener.in; filings from the exchanges; "
                "headlines from public news feeds. Verdicts and summaries are written by a language model running "
                "on this computer, using only that data. They can be wrong: this is not investment advice.",
           size=9, color=MUTED)

    for r in results:
        _w_company(doc, r)

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()

    # Same round-trip guarantee as the Excel export: fail loudly here rather
    # than silently shipping a .docx Word can't fully open.
    Document(io.BytesIO(data))
    return data


